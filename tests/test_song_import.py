import io
import zipfile

from pptx import Presentation
from song_import import import_batch, parse_song


def test_text_title_sections_and_attribution():
    raw=b'Hymn title\nAlso known as: Alternate title\n\n[Verse 1]\nFirst line\nSecond line\n\n[Chorus]\nSing together\n\n[Information]\nWords: Writer\nCopyright: permission notice'
    song=parse_song('001 - Hymn title.txt',raw)
    assert song['title']=='Hymn title'
    assert [s['label'] for s in song['slides']]==['Verse 1','Chorus']
    assert song['slides'][0]['text']=='First line\nSecond line'
    assert 'permission notice' in song['credit']
    assert 'Alternate title' in song['credit']


def test_zip_duplicates_and_invalid_member_are_isolated():
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as z:
        z.writestr('folder/Hymn.txt','Hymn\n\n[Verse 1]\nPraise together')
        z.writestr('copy/Hymn.txt','Hymn\n\n[Verse 1]\nPraise together')
        z.writestr('empty.txt','')
    songs,errors=import_batch([('collection.zip',buf.getvalue())])
    assert len(songs)==1
    assert len(errors)==1 and errors[0]['File']=='empty.txt'


def test_html_is_text_only_and_preserves_lines():
    raw=b'<html><head><title>My hymn</title><style>secret style</style></head><body><h1>My hymn</h1><p>[Verse 1]<br>Line one<br>Line two</p><script>alert(1)</script><p>[Chorus]<br>Praise</p></body></html>'
    song=parse_song('export.html',raw)
    assert song['title']=='My hymn'
    assert len(song['slides'])==2
    assert song['slides'][0]['text']=='Line one\nLine two'
    assert 'alert' not in str(song) and 'secret style' not in str(song)


def test_pptx_keeps_lyric_slides_and_legacy_ppt_reports_conversion():
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0,0,100,100).text='One short lyric line'
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(0,0,100,100).text='Chorus\nPraise together'
    buf=io.BytesIO();prs.save(buf)
    song=parse_song('Hymn.pptx',buf.getvalue())
    assert len(song['slides'])==2
    assert song['slides'][0]['text']=='One short lyric line'
    songs,errors=import_batch([('old.ppt',b'legacy')])
    assert not songs and '.pptx' in errors[0]['Reason']


def test_utf16_and_cp1252_are_decoded():
    assert parse_song('Hymn.txt','Hymn\n\n[Verse 1]\nGlory'.encode('utf-16'))['slides'][0]['text']=='Glory'
    assert 'Lord’s' in parse_song('Hymn.txt','Hymn\n\n[Verse 1]\nThe Lord’s love'.encode('cp1252'))['slides'][0]['text']


def test_title_only_is_skipped_and_different_credits_are_preserved():
    songs,errors=import_batch([('Title only.txt',b'Title only'),('a.txt',b'Hymn\n\n[Verse 1]\nPraise\n\n[Information]\nTune A'),('b.txt',b'Hymn\n\n[Verse 1]\nPraise\n\n[Information]\nTune B')])
    assert len(songs)==2 and len(errors)==1
    assert songs[0]['credit'] != songs[1]['credit']


def test_bundled_collection_excludes_questionable_sources_and_roundtrips():
    from song_import import load_bundled_collection,collection_review_issues
    from core import new_workspace,backup_bytes,load_backup
    bundle=load_bundled_collection()
    assert len(bundle['songs'])==315
    assert len(bundle['skipped'])==27
    assert all(not collection_review_issues(song) for song in bundle['songs'])
    workspace=new_workspace()
    workspace['library'].extend(bundle['songs'])
    restored=load_backup(backup_bytes(workspace))
    assert len(restored['library'])==318
