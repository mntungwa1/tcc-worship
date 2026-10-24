"""Import song text without executing HTML or extracting archive paths."""
import io
import base64
import gzip
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath
from pathlib import Path

from core import item

MAX_TOTAL = 40 * 1024 * 1024
MAX_FILE = 15 * 1024 * 1024
MAX_FILES = 900
SUPPORTED = {'.txt', '.html', '.htm', '.pptx', '.ppt'}


def load_bundled_collection():
    path = Path(__file__).with_name('data') / 'bundled_songs.json.gz.b64'
    return json.loads(gzip.decompress(base64.b64decode(path.read_text())))


def load_bundled_songs():
    return load_bundled_collection()['songs']


def collection_review_issues(song):
    """Conservative source-quality checks for the supplied labelled hymn collection."""
    issues = []
    verses = []
    for slide in song['slides']:
        match = re.fullmatch(r'Verse\s+(\d+)', slide['label'], re.I)
        if match:
            verses.append(int(match.group(1)))
        if re.search(r'^\s*\d+[.,]\s*\S', slide['text'], re.M):
            issues.append('Numbered verses embedded inside another verse; needs source review.')
        if re.search(r'^\s*(?:Words,?\s+Music|This hymn was written|Special thanks)', slide['text'], re.I | re.M):
            issues.append('Notes or unfinished attribution mixed into the lyrics.')
    if not verses or sorted(verses) != list(range(1, max(verses, default=0)+1)):
        issues.append('Verse numbering is missing, duplicated or has gaps.')
    if not song.get('credit'):
        issues.append('Attribution section is absent; source completeness needs review.')
    if not re.search(r'[.!?;:\"\u201d\u2019\x27]\s*$', song['slides'][-1]['text']):
        issues.append('Ending needs source review before inclusion.')
    return list(dict.fromkeys(issues))


def decode_text(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16')
    try:
        return raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        return raw.decode('cp1252')


class SongHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.title = []
        self.hidden = 0
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'head'):
            self.hidden += 1
        if tag == 'title':
            self.in_title = True
        if not self.hidden:
            if tag == 'br':
                self.parts.append('\n')
            elif tag in ('p', 'div', 'section', 'article', 'h1', 'h2', 'h3', 'li', 'tr', 'pre'):
                self.parts.append('\n\n')

    def handle_endtag(self, tag):
        if tag == 'title':
            self.in_title = False
        if tag in ('script', 'style', 'head'):
            self.hidden = max(0, self.hidden - 1)
        if not self.hidden and tag in ('p', 'div', 'section', 'article', 'h1', 'h2', 'h3', 'li', 'tr', 'pre'):
            self.parts.append('\n\n')

    def handle_data(self, data):
        if self.in_title:
            self.title.append(data)
        elif not self.hidden:
            self.parts.append(data)


def parse_text(text, filename, suggested_title=''):
    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\xa0', ' ').strip()
    fallback = re.sub(r'^\d+\s*[-_.]\s*', '', PurePosixPath(filename).stem)
    title = suggested_title.strip() or fallback
    if len(text.splitlines()) == 1 and (' '.join(text.casefold().split()) == ' '.join(title.casefold().split()) or all(label in text.casefold() for label in ('pdf image of score', 'gif image of score', 'midi audio of tune'))):
        raise ValueError('Title or website navigation only; no lyrics found.')
    credit = ''
    metadata = re.search(r'^\s*\[(?:Information|Copyright|Credits|Metadata)\]\s*$', text, re.I | re.M)
    if metadata:
        credit = text[metadata.end():].strip()
        text = text[:metadata.start()].strip()
    sections = list(re.finditer(r'^\s*\[([^\]\n]+)\][ \t]*$', text, re.M))
    slides = []
    if sections:
        preamble = text[:sections[0].start()].strip().splitlines()
        if preamble:
            title = preamble[0].strip()
            extra = '\n'.join(preamble[1:]).strip()
            if extra:
                credit = '\n'.join(filter(None, [extra, credit]))
        for i, match in enumerate(sections):
            body = text[match.end():sections[i+1].start() if i+1 < len(sections) else len(text)].strip()
            if body:
                slides.append({'label': match.group(1).strip(), 'text': body})
    else:
        lines = text.splitlines()
        # Only strip an unlabelled title when separated from the lyrics, or matching the filename/HTML title.
        if len(lines) > 1 and (lines[0].strip().casefold() in (fallback.casefold(), title.casefold()) or (not lines[1].strip() and len(lines[0]) < 150)):
            title = lines.pop(0).strip()
            text = '\n'.join(lines).strip()
        slides = item(title, 'song', text)['slides']
    if not slides:
        raise ValueError('No lyrics found.')
    if len(slides) > 500:
        raise ValueError('More than 500 lyric slides; split this file into songs.')
    return item(title[:300], 'song', credit=credit, slides=slides)


def parse_song(filename, raw):
    suffix = PurePosixPath(filename).suffix.lower()
    if len(raw) > MAX_FILE:
        raise ValueError('File exceeds 15 MB.')
    if suffix == '.ppt':
        raise ValueError('Older .ppt format: open in PowerPoint or LibreOffice, Save As .pptx, then import again.')
    if suffix == '.pptx':
        with zipfile.ZipFile(io.BytesIO(raw)) as package:
            if sum(i.file_size for i in package.infolist()) > MAX_TOTAL:
                raise ValueError('PowerPoint expanded content exceeds 40 MB.')
        from pptx import Presentation
        prs = Presentation(io.BytesIO(raw))
        slides = []
        for n, slide in enumerate(prs.slides, 1):
            text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame).strip()
            if text:
                slides.append({'label': f'Slide {n}', 'text': text})
        if not slides:
            raise ValueError('No editable text found in PowerPoint; image-only slides cannot become song lyrics.')
        if len(slides) > 500:
            raise ValueError('More than 500 slides; split this presentation.')
        title = re.sub(r'^\d+\s*[-_.]\s*', '', PurePosixPath(filename).stem)
        # Title-only slides are metadata; keep every lyric slide and its line breaks.
        if len(slides) > 1 and slides[0]['text'].casefold() == title.casefold():
            title = slides.pop(0)['text']
        return item(title[:300], 'song', slides=slides)
    if suffix in ('.html', '.htm'):
        parser = SongHTML()
        parser.feed(decode_text(raw))
        return parse_text(''.join(parser.parts), filename, ''.join(parser.title))
    if suffix == '.txt':
        return parse_text(decode_text(raw), filename)
    raise ValueError('Unsupported song format.')


def fingerprint(song):
    normal = lambda value: ' '.join(value.casefold().split())
    return (normal(song['title']), tuple((normal(s.get('label','')), normal(s['text'])) for s in song['slides']), normal(song.get('credit','')))


def import_batch(files):
    """Return valid songs and individual errors; never extract archive members to disk."""
    songs, errors, seen = [], [], set()
    total = 0
    count = 0

    def consume(name, raw):
        nonlocal total, count
        count += 1
        total += len(raw)
        if count > MAX_FILES or total > MAX_TOTAL:
            raise ValueError('Batch exceeds 900 files or 40 MB expanded text/media. Split the collection.')
        try:
            song = parse_song(name, raw)
            key = fingerprint(song)
            if key not in seen:
                song['import_source'] = name
                songs.append(song)
                seen.add(key)
        except Exception as exc:
            errors.append({'File': name, 'Reason': str(exc)})

    for name, raw in files:
        if len(raw) > MAX_TOTAL:
            errors.append({'File': name, 'Reason': 'Upload exceeds 40 MB.'})
            continue
        if PurePosixPath(name).suffix.lower() != '.zip':
            consume(name, raw)
            continue
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                members = [i for i in archive.infolist() if not i.is_dir() and not i.filename.startswith('__MACOSX/')]
                if len(members) > MAX_FILES or sum(i.file_size for i in members) > MAX_TOTAL:
                    raise ValueError('ZIP exceeds 900 files or 40 MB expanded. Split the archive.')
                supported = [i for i in members if PurePosixPath(i.filename).suffix.lower() in SUPPORTED]
                if not supported:
                    raise ValueError('No TXT, HTML or PowerPoint songs found inside this ZIP.')
                for member in supported:
                    if member.file_size > MAX_FILE:
                        errors.append({'File': member.filename, 'Reason': 'File exceeds 15 MB.'})
                        continue
                    consume(member.filename, archive.read(member))
        except ValueError:
            raise
        except Exception as exc:
            errors.append({'File': name, 'Reason': 'Could not read ZIP: ' + str(exc)})
    return songs, errors
