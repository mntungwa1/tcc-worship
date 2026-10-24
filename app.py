import copy
import hmac
import html
import io
import json
import os
from pathlib import Path

import streamlit as st
import streamlit.components.v1 as components
from core import (new_workspace, item, duplicate, slides_from_text, search_items, presentation_item,
                  backup_bytes, load_backup, media_data, export_html, MAX_BACKUP, FONT_FAMILIES)
from bible import BOOKS, load_bible, validate_bible, parse_reference, passage_item
from song_import import load_bundled_songs, fingerprint
from saved_services import controls as saved_service_controls, restore_project
from media_catalog import (CATEGORIES, asset_category, import_uploads, preview_source,
                           embed_source, portable_item, archive_library)

st.set_page_config(page_title='TCC Worship | Tabernacle Christian Centre', page_icon='✝', layout='wide')
st.markdown('''<style>
.block-container{padding-top:2rem}h1{letter-spacing:-.04em}div[data-testid="stMetric"]{background:#142039;border:1px solid #2c3a53;border-radius:12px;padding:14px}div[data-testid="stSidebar"]{border-right:1px solid #27354d}.eyebrow{color:#e8bc73;font-size:12px;letter-spacing:2px}.subtle{color:#a7b6cc}.stButton>button{border-radius:8px}.desk-title{font-size:22px;font-weight:650}.bible-label{font-size:12px;color:#e8bc73;letter-spacing:1px;margin-bottom:8px}

/* Keep the first label clear of Streamlit's fixed toolbar. */
.block-container{padding-top:6rem}
div[data-testid="stMetric"]{color:#f4f6fb}
div[data-testid="stMetric"] [data-testid="stMetricLabel"],
div[data-testid="stMetric"] [data-testid="stMetricValue"],
div[data-testid="stMetric"] p{color:#f4f6fb}
div[data-testid="stMetricLabel"] p,
div[data-testid="stWidgetLabel"] p,
div[data-testid="stRadio"] label p,
div[data-testid="stButton"] button p,
div[data-testid="stDownloadButton"] button p,
div[data-testid="stFormSubmitButton"] button p{
    white-space:normal;overflow:visible;text-overflow:clip;overflow-wrap:anywhere
}
div[data-testid="stButton"] button,
div[data-testid="stDownloadButton"] button,
div[data-testid="stFormSubmitButton"] button{height:auto;min-height:2.5rem}
.eyebrow{line-height:1.6;overflow-wrap:anywhere}

</style>''', unsafe_allow_html=True)

def secret(name, default=''):
    try: return st.secrets.get(name, os.getenv(name, default))
    except Exception: return os.getenv(name, default)

password = secret('APP_PASSWORD')
if password and not st.session_state.get('authenticated'):
    st.image(str(Path(__file__).parent/'assets/tcc_logo.png'), width=340)
    st.title('TCC Worship')
    st.caption('TABERNACLE CHRISTIAN CENTRE • WORSHIP DESK')
    with st.form('login'):
        entered=st.text_input('Church access password',type='password')
        if st.form_submit_button('Open worship desk',type='primary'):
            if hmac.compare_digest(entered,password):
                st.session_state.authenticated=True;st.rerun()
            else:st.error('Incorrect password.')
    st.stop()

if 'workspace' not in st.session_state:
    st.session_state.workspace = new_workspace()
    st.session_state.bibles = [load_bible()]
    st.session_state.verse_history = []
    st.session_state.verse_favourites = []
pending=st.session_state.pop('pending_saved_service',None)
if pending:
    try:
        if pending['action']=='new':
            st.session_state.workspace['service']=[]
            st.session_state.workspace['service_title']='New Worship Service'
            for key in ('saved_service_id','saved_service_name','saved_service_revision'):
                st.session_state.pop(key,None)
            st.session_state.notice='New service started. Your song and media libraries are ready to use.'
        else:
            workspace,bibles=restore_project(st.session_state.workspace,pending['project'])
            st.session_state.workspace=workspace
            st.session_state.bibles=bibles
            st.session_state.saved_service_id=pending.get('id')
            st.session_state.saved_service_name=pending.get('name') or workspace['service_title']
            st.session_state.saved_service_revision=pending.get('revision')
            st.session_state.notice='Opened '+st.session_state.saved_service_name+'.'
        st.session_state.pop('live_desk_html',None)
        st.session_state.pop('live_desk_signature',None)
    except (ValueError,KeyError,TypeError) as exc:
        st.session_state.notice='Could not open the saved preparation: '+str(exc)
w=st.session_state.workspace
if 'presenter_session' not in st.session_state:
    st.session_state.presenter_session=os.urandom(12).hex()

# Seed the supplied collection once per session, including sessions opened before this update.
if not st.session_state.get('bundled_songs_v1_loaded'):
    existing_songs={fingerprint(x) for x in w['library'] if x['kind']=='song'}
    additions=[x for x in load_bundled_songs() if fingerprint(x) not in existing_songs]
    candidate=copy.deepcopy(w)
    candidate['library'].extend(additions)
    if len(candidate['library'])<=1000 and len(backup_bytes(candidate))<=MAX_BACKUP:
        w['library']=candidate['library']
        st.session_state.bundled_songs_v1_loaded=True
    else:
        st.warning('The included song collection could not fit in this workspace. Save a backup and start a fresh session to load it.')



def notify(message):
    st.session_state.notice=message
    st.rerun()

def queue(obj):
    try:
        candidate=copy.deepcopy(w)
        candidate['service'].append(duplicate(portable_item(obj)))
        if len(backup_bytes(candidate))>MAX_BACKUP:
            raise ValueError('This prepared service would exceed 45 MB. Choose smaller media or split the service.')
        w['service']=candidate['service']
        notify('Added to the service.')
    except ValueError as exc:st.error(str(exc))

def preview(obj, slide=0, height=260, show_background=False):
    if obj['kind']=='song':
        obj=presentation_item(obj)
        slide=min(slide,len(obj['slides'])-1)
        st.caption(f"{obj['slides'][slide]['label']} · Slide {slide+1} of {len(obj['slides'])}")
    if obj['kind'] in ('image','video'):
        try:
            src=preview_source(obj['src'])
            if obj['kind']=='video':st.video(src)
            else:st.image(src)
        except ValueError as exc:st.info(str(exc))
        return
    s=obj['slides'][slide]
    if s.get('src'):st.image(s['src']);return
    text=html.escape(s.get('text','')).replace('\n','<br>')
    credit=html.escape(obj.get('credit',''))
    backdrop=w['theme'].get('background_media') if show_background else None
    layer=''
    if backdrop:
        source=html.escape(backdrop['src'],quote=True)
        style='position:absolute;inset:0;width:100%;height:100%;object-fit:cover;pointer-events:none'
        layer=(f'<video src="{source}" style="{style}" muted autoplay loop playsinline></video>' if backdrop['kind']=='video' else f'<img src="{source}" alt="" style="{style}">')
        layer+=f'<div style="position:absolute;inset:0;background:black;opacity:{w["theme"].get("background_overlay",0.45)}"></div>'
    font=html.escape(w['theme'].get('font_family','Arial'),quote=True)
    size=w['theme']['font_size']
    components.html(f'''<div style="position:relative;box-sizing:border-box;height:{height-10}px;background:{w['theme']['background']};color:white;font-family:'{font}';font-size:clamp(14px,{size/12.8}vw,{size}px);line-height:1.35;display:flex;flex-direction:column;justify-content:center;align-items:center;text-align:center;border:1px solid #34435c;border-radius:10px;padding:20px;overflow:hidden">{layer}<div dir="auto" style="position:relative;z-index:2;overflow-wrap:anywhere">{text}</div><div style="position:relative;z-index:2;font-size:11px;color:#d4d9e2;margin-top:16px">{credit}</div></div>''',height=height)


def set_background(asset,overlay=None):
    candidate=copy.deepcopy(w)
    if asset is None:
        candidate['theme'].pop('background_media',None)
    else:
        candidate['theme']['background_media']={'kind':asset['kind'],'src':embed_source(asset['src']),'name':asset['title'],'library_id':asset['id'],'category':asset_category(asset)}
    if overlay is not None:candidate['theme']['background_overlay']=overlay
    if len(backup_bytes(candidate))>MAX_BACKUP:
        raise ValueError('The workspace would exceed 45 MB. Use smaller media or remove unused items.')
    w['theme']=candidate['theme']


def background_controls(prefix):
    pending=st.session_state.pop(prefix+'_pending_bg_choice',None)
    if pending:
        st.session_state[prefix+'_bg_type']=pending['mode']
        st.session_state[prefix+'_bg_category_'+pending['mode']]=pending['category']
        st.session_state[prefix+'_bg_'+pending['mode']]=pending['id']
    with st.expander('Image or video background'):
        current=w['theme'].get('background_media')
        options=['Solid colour','Image','Video']
        default=options.index(current['kind'].title()) if current else 0
        mode=st.radio('Background type',options,index=default,horizontal=True,key=prefix+'_bg_type')
        chosen=None
        if mode!='Solid colour':
            kind=mode.lower()
            all_assets=[x for x in w['library'] if x['kind']==kind]
            counts={name:sum(asset_category(x)==name for x in all_assets) for name in CATEGORIES}
            category=st.selectbox('Background category',['All categories',*CATEGORIES],format_func=lambda name:f"{name} ({len(all_assets) if name=='All categories' else counts[name]})",key=prefix+'_bg_category_'+mode)
            assets=[x for x in all_assets if category=='All categories' or asset_category(x)==category]
            ids=[x['id'] for x in assets]
            default_id=current.get('library_id') if current and current['kind']==kind else None
            selected=st.selectbox('Choose an '+kind+' file' if kind=='image' else 'Choose a video file',[None,*ids],index=ids.index(default_id)+1 if default_id in ids else 0,format_func=lambda value:('Select a file from the list' if ids else 'No files in this category yet') if value is None else next(x['title'] for x in assets if x['id']==value),disabled=not ids,key=prefix+'_bg_'+mode)
            chosen=next((x for x in assets if x['id']==selected),None)
            if chosen:
                st.caption('Selected file: '+chosen['title'])
                preview(chosen,height=220)
            elif not assets:
                if all_assets:
                    st.info(f'No {kind} files in this category. Choose All categories to see your {len(all_assets)} uploaded files, or browse your computer below.')
                else:
                    st.info('No '+kind+' files uploaded yet. Click Browse files below to select files from your computer.')
            st.markdown('**Add a background from your computer**')
            st.caption('Click Browse files to open File Explorer. Select one or more files, or an Images.zip / Videos.zip collection, then click Add selected files below.')
            upload_category=st.selectbox('Category for new files',CATEGORIES,index=CATEGORIES.index(category) if category in CATEGORIES else CATEGORIES.index('Welcome'),key=prefix+'_bg_upload_category_'+mode)
            version=st.session_state.get(prefix+'_bg_upload_version_'+mode,0)
            extensions=['png','jpg','jpeg','webp','gif','zip'] if kind=='image' else ['mp4','webm','zip']
            uploads=st.file_uploader('Browse your computer for '+kind+' files or ZIP',type=extensions,accept_multiple_files=True,key=prefix+'_bg_upload_'+mode+'_'+str(version))
            if uploads:
                st.write('Selected files to add:')
                st.dataframe([{'File':file.name,'Size (MB)':round(file.size/1024/1024,2)} for file in uploads],hide_index=True,width='stretch')
            if st.button('Add selected files to background list',disabled=not uploads,key=prefix+'_bg_upload_add_'+mode):
                try:
                    with st.spinner('Adding your background files…'):
                        result=import_uploads(w,uploads,kind,upload_category,True)
                    if result['asset_ids']:
                        first=next(x for x in w['library'] if x['id']==result['asset_ids'][0])
                        st.session_state[prefix+'_pending_bg_choice']={'mode':mode,'category':asset_category(first),'id':first['id']}
                        st.session_state[prefix+'_bg_upload_version_'+mode]=version+1
                        st.session_state.pop(kind+'_prepared_archive',None)
                        notify(f"Added {result['added']} new files; restored {result['restored']}. The first file is selected in the list. Click Apply background when ready.")
                    else:st.warning('No matching '+kind+' files were found. Choose supported '+kind+' files or the matching media ZIP.')
                except (ValueError,OSError) as exc:st.error(str(exc))
        darkening=st.slider('Darken background for readable lyrics (%)',0,90,int(w['theme'].get('background_overlay',0.45)*100),key=prefix+'_bg_dark')
        if st.button('Apply background',disabled=mode!='Solid colour' and chosen is None,key=prefix+'_bg_apply',type='primary'):
            try:
                set_background(chosen,darkening/100)
                notify('Background updated. In Presentation desk, click Load / refresh live desk when ready.')
            except ValueError as exc:st.error(str(exc))
        if current:st.caption('Current background: '+current['name'])
        st.caption('Applies behind lyrics, scripture and text slides. Video backgrounds loop silently. Image and video service items use their own content.')
        if prefix=='settings':preview(item('Background preview','announcement','Background preview\nLet us worship together.'),height=220,show_background=True)


def media_library(kind):
    extensions=['png','jpg','jpeg','webp','gif'] if kind=='image' else ['mp4','webm']
    default=st.selectbox('Category for loose files',CATEGORIES,index=CATEGORIES.index('Welcome'),key=kind+'_upload_category')
    folders=st.checkbox('Use category folders in ZIP files',value=True,key=kind+'_use_folders')
    st.caption('ZIP folders such as Praise, Communion or Sunday School assign categories automatically. Files without a matching folder use the category selected above.')
    version=st.session_state.get(kind+'_upload_version',0)
    uploads=st.file_uploader('Add '+kind+' files or ZIP collections',type=extensions+['zip'],accept_multiple_files=True,key=kind+'_library_uploads_'+str(version))
    st.caption('ZIP upload: up to 512 MB. Each media file: up to 20 MB. Your full image/video collection can hold up to 768 MB.')
    if st.button('Add files to '+kind+' library',disabled=not uploads,type='primary',key=kind+'_library_add'):
        try:
            with st.spinner('Importing files into their categories…'):
                result=import_uploads(w,uploads,kind,default,folders)
            st.session_state[kind+'_upload_version']=version+1
            st.session_state.pop(kind+'_prepared_archive',None)
            detail=', '.join(f'{name}: {count}' for name,count in result['categories'].items() if count)
            message=f"Added {result['added']} files · restored {result['restored']} · duplicates skipped {result['duplicates']}."
            if detail:message+=' '+detail+'.'
            if result['ignored']:message+=f" Ignored {result['ignored']} supporting files or files for the other library."
            notify(message)
        except (ValueError,OSError) as exc:st.error(str(exc))
    all_assets=[x for x in w['library'] if x['kind']==kind]
    counts={name:sum(asset_category(x)==name for x in all_assets) for name in CATEGORIES}
    st.caption(f'{len(all_assets)} items in your {kind} library · 10 categories')
    category=st.selectbox('Browse category',['All categories',*CATEGORIES],format_func=lambda name:f"{name} ({len(all_assets) if name=='All categories' else counts[name]})",key=kind+'_browse_category')
    query=st.text_input('Search '+kind+' library',key=kind+'_library_search')
    assets=[x for x in all_assets if (category=='All categories' or asset_category(x)==category) and query.casefold() in (x['title']+' '+x.get('credit','')).casefold()]
    if not assets:st.info('No matching media in this category yet. Upload files above or choose another category.')
    pages=max(1,(len(assets)+11)//12)
    page=st.selectbox('Page',range(1,pages+1),key=kind+'_media_page')
    for obj in assets[(page-1)*12:page*12]:
        with st.expander(obj['title']+' · '+asset_category(obj)):
            preview(obj)
            if obj.get('credit'):
                with st.expander('Source and licence credits'):st.text(obj['credit'])
            new_category=st.selectbox('Move to category',CATEGORIES,index=CATEGORIES.index(asset_category(obj)),key='media_category_'+obj['id'])
            if st.button('Save category',key='media_category_save_'+obj['id']):
                obj['category']=new_category
                st.session_state.pop(kind+'_prepared_archive',None)
                notify('Moved to '+new_category+'.')
            a,b=st.columns(2)
            if a.button('Add to service',key='media_queue_'+obj['id']):queue(obj)
            if b.button('Use as background',key='media_background_'+obj['id']):
                try:
                    set_background(obj)
                    notify(obj['title']+' selected as background. Load / refresh the live desk when ready.')
                except ValueError as exc:st.error(str(exc))
    archive_signature=json.dumps([(x['id'],x['src'],asset_category(x)) for x in all_assets])
    if st.session_state.get(kind+'_archive_signature')!=archive_signature:
        st.session_state.pop(kind+'_prepared_archive',None)
    if all_assets:
        with st.expander('Keep a copy of this '+kind+' collection'):
            st.caption('Keep the original ZIPs, or prepare a new category-organised ZIP here. Upload it again if a server restart clears the media files.')
            if st.button('Prepare '+kind+' library ZIP',key=kind+'_prepare_archive'):
                try:
                    with st.spinner('Preparing your media ZIP…'):
                        st.session_state[kind+'_prepared_archive']=archive_library(w,kind)
                        st.session_state[kind+'_archive_signature']=archive_signature
                except (ValueError,OSError) as exc:st.error(str(exc))
            packed=st.session_state.get(kind+'_prepared_archive')
            if packed:
                st.download_button('Download '+kind+' library ZIP',packed,file_name='Images.zip' if kind=='image' else 'Videos.zip',mime='application/zip',key=kind+'_archive_download')


@st.cache_data(ttl=86400, max_entries=12, show_spinner=False)
def cached_bible_version(version_id):
    from bible_catalogue import fetch_translation
    return fetch_translation(version_id)

def bible_catalogue_controls():
    from bible_catalogue import load_catalogue
    catalogue=load_catalogue()
    with st.expander(f'Browse {len(catalogue)} free Bible versions'):
        st.caption('Twelve selected English translations. Add the versions you use; internet is needed for the first download. Save Bible JSON to retain a version across sessions, or include it in an offline presentation.')
        purposes={'BSB':'Modern English for everyday reading and projection.','eng_kjv':'Traditional English.','eng_web':'Modern English, public domain.','eng_asv':'Classic formal translation.','eng_bbe':'Simpler English vocabulary.','eng_fbv':'Readable contemporary English.','eng_dby':'A close, literal translation.','eng_ylt':'A very literal study translation.','eng_rv5':'Historic revision for comparison.','eng_wbs':"Webster’s revision of traditional English wording.",'eng_gnv':'Historic Reformation translation.','eng_lsv':'Modern literal translation.'}
        selected=st.selectbox('Add a Bible version',range(len(catalogue)),format_func=lambda i:f"{catalogue[i]['englishName']} · {catalogue[i]['shortName']}",key='catalogue_version')
        entry=catalogue[selected]
        st.write(purposes.get(entry['id'],''))
        rights=entry['license']
        st.caption(('Public domain' if rights['publicDomain'] else 'Open licence — attribution retained')+f" · {entry['numberOfBooks']} books")
        st.link_button('View source and licence',entry['licenseUrl'])
        present=any(x.get('id')==entry['id'] or (entry['id']=='eng_web' and not x.get('id') and x['name'].startswith('World English Bible')) for x in st.session_state.bibles)
        if st.button('Add version to Bible desk',disabled=present,type='primary',key='catalogue_add'):
            try:
                with st.spinner('Downloading and checking the selected Bible version…'):
                    downloaded=cached_bible_version(entry['id'])
                st.session_state.bibles.append(downloaded)
                notify(downloaded['name']+' added. Select it in the Bible version dropdown.')
            except Exception as exc:
                st.error('Could not add this version: '+str(exc))
        if present:
            st.success('This version is already available in the Bible version dropdown.')

with st.sidebar:
    st.image(str(Path(__file__).parent/'assets/tcc_logo.png'),width='stretch')
    st.markdown('<div class="eyebrow">TABERNACLE CHRISTIAN CENTRE</div>',unsafe_allow_html=True)
    st.title('TCC Worship')
    st.caption('Prepare with purpose. Present with confidence.')
    page=st.radio('Workspace',['Overview','Songs & library','Bible desk','Service planner','Media & presentations','Presentation desk','Saved services','Settings & backups'],label_visibility='collapsed')
    st.divider()
    st.caption(f"{len(w['library'])} library items · {len(w['service'])} service items")
    st.info('Use Saved services to keep named preparations in this browser. Download service files or workspace backups to keep a separate copy.')
    if password and st.button('Lock desk'):
        st.session_state.authenticated=False;st.rerun()

st.markdown('<div class="eyebrow">TCC WORSHIP / TABERNACLE CHRISTIAN CENTRE</div>',unsafe_allow_html=True)
st.title(page)
if st.session_state.get('notice'):st.success(st.session_state.pop('notice'))

if page=='Overview':
    st.image(str(Path(__file__).parent/'assets/tcc_logo.png'),width=420)
    st.markdown('### One place to prepare your worship service')
    st.write('Build the service, choose your songs and scripture, then open a clean projector and stage view from your offline presentation package.')
    a,b,c=st.columns(3)
    a.metric('Library',len(w['library']))
    b.metric('Order of service',len(w['service']))
    c.metric('Bible versions',len(st.session_state.bibles))
    st.subheader(w['service_title'])
    for n,obj in enumerate(w['service'],1):st.write(f"**{n:02}. {obj['title']}** · {obj['kind'].title()}")
    st.caption('Start in Songs & library or Bible desk. Add items to Service planner, then use Presentation desk to export and present.')

elif page=='Songs & library':
    browse,edit,bulk=st.tabs(['Find & select','Create / edit','Bulk import'])
    with browse:
        cols=st.columns([3,2,1])
        q=cols[0].text_input('Search titles, lyrics and tags')
        lang=cols[1].selectbox('Language',['All']+sorted({x.get('language','') for x in w['library'] if x.get('language')}))
        fav=cols[2].checkbox('Favourites only')
        found=search_items(w['library'],q,lang,fav)
        if not found:st.info('No matching items. Create one in the next tab.')
        for obj in found:
            with st.expander(f"{'★ ' if obj.get('favourite') else ''}{obj['title']} · {obj['kind']}"):
                a,b=st.columns([3,2])
                with a:preview(obj)
                with b:
                    st.caption(obj.get('tags',''))
                    if st.button('Add to service',key='add'+obj['id'],type='primary'):queue(obj)
                    if st.button('Toggle favourite',key='fav'+obj['id']):obj['favourite']=not obj.get('favourite');notify('Favourite updated.')
    with edit:
        editable=[x for x in w['library'] if x['kind'] in ('song','announcement','scripture')]
        selected=st.selectbox('Edit an existing item',['New item']+[x['id'] for x in editable],format_func=lambda v: 'New item' if v=='New item' else next(x['title'] for x in editable if x['id']==v))
        original=next((x for x in editable if x['id']==selected),None)
        with st.form('editor'+selected):
            title=st.text_input('Title',value=original['title'] if original else '')
            kind=st.selectbox('Type',['song','announcement','scripture'],index=['song','announcement','scripture'].index(original['kind']) if original else 0)
            language=st.text_input('Language',value=original.get('language','') if original else '')
            tags=st.text_input('Tags / themes',value=original.get('tags','') if original else '')
            txt='\n\n'.join((('['+s['label']+']\n') if s.get('label') else '')+s['text'] for s in original['slides']) if original else ''
            text=st.text_area('Lyrics or text — separate slides with a blank line; use [Chorus] for labels',value=txt,height=250)
            credit=st.text_input('Author / copyright / licence attribution',value=original.get('credit','') if original else '')
            if st.form_submit_button('Save to library',type='primary'):
                if not title.strip() or not slides_from_text(text):st.error('Add a title and at least one slide of text.')
                else:
                    obj=item(title.strip(),kind,text,language=language,tags=tags,credit=credit)
                    if original:obj['id']=original['id'];obj['favourite']=original.get('favourite',False);w['library'][w['library'].index(original)]=obj
                    else:w['library'].append(obj)
                    notify('Saved. Existing service copies stay as prepared; add the revised item to use it.')


    with bulk:
        st.write('Import a collection of songs from TXT, HTML or PowerPoint files. Select multiple files, or upload a ZIP containing them.')
        st.caption('One file becomes one song. Verse/chorus labels and text-file credits are retained. PPTX imports editable text only; save older .ppt files as .pptx first. Review HTML exports for navigation text.')
        from song_import import import_batch, fingerprint, load_bundled_collection
        included=load_bundled_collection()
        st.caption(f"Included hymn collection: {len(included['songs'])} songs. {len(included['skipped'])} source files held out for review.")
        with st.expander('Files excluded from the included collection'):
            st.dataframe(included['skipped'],hide_index=True,width='stretch')
        uploads=st.file_uploader('Song files or ZIP collections',type=['txt','html','htm','pptx','ppt','zip'],accept_multiple_files=True,key='song_uploads')
        if st.button('Read files and preview songs',disabled=not uploads,type='primary'):
            try:
                if sum(f.size for f in uploads)>40*1024*1024:
                    raise ValueError('Select at most 40 MB per batch.')
                songs,errors=import_batch([(f.name,f.getvalue()) for f in uploads])
                st.session_state.song_import_preview={'songs':songs,'errors':errors}
            except Exception as exc:
                st.session_state.pop('song_import_preview',None)
                st.error(str(exc))
        batch=st.session_state.get('song_import_preview')
        if batch is not None:
            songs=batch['songs']
            existing={fingerprint(x) for x in w['library'] if x['kind']=='song'}
            pending=[x for x in songs if fingerprint(x) not in existing]
            st.info(f"{len(songs)} unique songs read · {len(pending)} new · {len(songs)-len(pending)} already in your library.")
            if batch['errors']:
                st.warning(f"{len(batch['errors'])} files could not be imported. Details below.")
                st.dataframe(batch['errors'],hide_index=True,width='stretch')
            if songs:
                st.dataframe([{'Title':x['title'],'Slides':len(x['slides']),'Source':x.get('import_source',''),'Status':'Already in library' if fingerprint(x) in existing else 'New'} for x in songs],hide_index=True,width='stretch')
                sample=st.selectbox('Preview an imported song',range(len(songs)),format_func=lambda i:songs[i]['title'],key='import_song_preview')
                preview(songs[sample])
                if songs[sample].get('credit'):
                    st.caption(songs[sample]['credit'])
                confirm_import=st.checkbox('I have reviewed the preview and want to add all new songs',key='confirm_song_import')
                if st.button(f'Import {len(pending)} new songs to library',disabled=not(pending and confirm_import),type='primary'):
                    try:
                        candidate=copy.deepcopy(w)
                        candidate['library'].extend(copy.deepcopy(pending))
                        if len(backup_bytes(candidate))>MAX_BACKUP:
                            raise ValueError('The workspace would exceed 45 MB. Import a smaller batch.')
                        w['library']=candidate['library']
                        notify(f'Imported {len(pending)} songs. Download a workspace backup in Settings & backups to keep them.')
                    except ValueError as exc:
                        st.error(str(exc))
            st.caption('Imports stay in this session. Download a workspace backup after importing, and restore it next time.')

elif page=='Bible desk':
    st.caption('Holyrics-inspired Bible workspace • verse list on the left • book, chapter and verse grids on the right')
    bible_catalogue_controls()
    versions=st.session_state.bibles
    vi=st.selectbox('Bible version',range(len(versions)),format_func=lambda i:versions[i]['name'])
    bible=versions[vi]
    if bible.get('license_notice'):
        with st.expander('Translation attribution and licence'):
            st.text(bible['license_notice'])
    st.download_button('Download selected Bible JSON',json.dumps(bible,ensure_ascii=False).encode('utf-8'),file_name='TCC_Bible_'+str(bible.get('id','WEB'))+'.json',mime='application/json')
    with st.form('reference'):
        a,b=st.columns([5,1])
        ref=a.text_input('Quick reference',placeholder='John 3:16-18 or Psalms 23:1-6')
        go=b.form_submit_button('Find passage',type='primary')
    if go:
        try:
            book,ch,nums=parse_reference(ref,bible)
            st.session_state.bible_book=book;st.session_state.bible_chapter=int(ch)
            st.session_state.passage_selection=nums
        except ValueError as exc:st.error(str(exc))
    books=[x for x in BOOKS if x in bible['books']]+[x for x in bible['books'] if x not in BOOKS]
    if st.session_state.get('bible_book') not in books:st.session_state.bible_book=books[0]
    book=st.session_state.bible_book
    chapters=sorted(map(int,bible['books'][book]))
    if st.session_state.get('bible_chapter') not in chapters:st.session_state.bible_chapter=chapters[0]
    chapter=st.session_state.bible_chapter
    verses=bible['books'][book][str(chapter)]
    left,right=st.columns([1.25,1.05],gap='large')
    with right:
        st.markdown('<div class="bible-label">BOOKS</div>',unsafe_allow_html=True)
        group=st.radio('Testament',['Old Testament','New Testament','All'],horizontal=True,label_visibility='collapsed')
        visible=[b for b in books if group=='All' or (b in BOOKS[:39] if group=='Old Testament' else b in BOOKS[39:])]
        grid=st.columns(5)
        for n,b in enumerate(visible):
            if grid[n%5].button(b,key='book_'+b,type='primary' if b==book else 'secondary',width='stretch'):
                st.session_state.bible_book=b;st.session_state.bible_chapter=min(map(int,bible['books'][b]));st.session_state.passage_selection=[];st.rerun()
        st.markdown('<div class="bible-label">CHAPTERS</div>',unsafe_allow_html=True)
        grid=st.columns(10)
        for n,ch in enumerate(chapters):
            if grid[n%10].button(str(ch),key=f'ch_{book}_{ch}',type='primary' if ch==chapter else 'secondary',width='stretch'):
                st.session_state.bible_chapter=ch;st.session_state.passage_selection=[];st.rerun()
        st.markdown('<div class="bible-label">VERSES</div>',unsafe_allow_html=True)
        grid=st.columns(10)
        for n,v in enumerate(sorted(map(int,verses))):
            if grid[n%10].button(str(v),key=f'v_{book}_{chapter}_{v}',width='stretch'):
                st.session_state.passage_selection=[v];st.rerun()
    with left:
        st.subheader(f'{book} {chapter}')
        verse_numbers=sorted(map(int,verses))
        selection=[n for n in st.session_state.get('passage_selection',[]) if n in verse_numbers] or [verse_numbers[0]]
        # A changing widget key prevents a previous chapter's selection from leaking into this one.
        choose=st.multiselect('Select verses to present',verse_numbers,default=selection,key=f'sel_{vi}_{book}_{chapter}_{selection}')
        with st.container(height=260,border=True):
            for n in verse_numbers:
                color='#e8bc73' if n in choose else '#b9c7d9'
                st.markdown(f'<p style="color:{color}"><strong>{n}</strong> {html.escape(verses[str(n)])}</p>',unsafe_allow_html=True)
        if choose:
            passage=passage_item(bible,book,chapter,choose)
            preview(passage,height=240)
            a,b=st.columns(2)
            if a.button('Add passage to service',type='primary',width='stretch'):
                st.session_state.verse_history.insert(0,copy.deepcopy(passage));queue(passage)
            if b.button('★ Save passage',width='stretch'):
                st.session_state.verse_favourites.append(copy.deepcopy(passage));notify('Passage saved for this session.')
            comparison=st.checkbox('Compare with another imported version')
            if comparison:
                other=[x for i,x in enumerate(versions) if i!=vi]
                if not other:st.info('Import another authorised Bible version in Settings & backups.')
                else:
                    oi=st.selectbox('Compare version',range(len(other)),format_func=lambda i:other[i]['name'])
                    for n in choose:st.write(f"{n}. {other[oi]['books'].get(book,{}).get(str(chapter),{}).get(str(n),'Verse unavailable in this version.')}")
    st.divider()
    search=st.text_input('Search words or phrases throughout this Bible')
    if search.strip():
        hits=[]
        for bk,chs in bible['books'].items():
            for ch,vs in chs.items():
                for n,text in vs.items():
                    if search.casefold() in text.casefold():hits.append((bk,ch,n,text))
        st.caption(f'{len(hits)} matches · first 40 shown')
        for bk,ch,n,text in hits[:40]:
            if st.button(f'{bk} {ch}:{n} — {text[:110]}',key=f'hit_{bk}_{ch}_{n}'):
                st.session_state.bible_book=bk;st.session_state.bible_chapter=int(ch);st.session_state.passage_selection=[int(n)];st.rerun()
    for label,key in [('Favourite passages','verse_favourites'),('Passage history','verse_history')]:
        with st.expander(label):
            for n,obj in enumerate(st.session_state[key][:30]):
                if st.button('Add '+obj['title'],key=f'{key}_{n}'):queue(obj)

elif page=='Service planner':
    w['service_title']=st.text_input('Service title',value=w['service_title']) or 'TCC Worship Service'
    st.caption('Use the arrows to reorder. Items are independent copies, keeping your prepared service stable when library entries change.')
    if not w['service']:st.info('Add songs, scripture or media to begin.')
    for n,obj in enumerate(w['service']):
        with st.container(border=True):
            a,b,c,d=st.columns([8,1,1,1])
            a.write(f"**{n+1:02}. {obj['title']}**")
            a.caption(f"{obj['kind'].title()} · {max(1,len(presentation_item(obj)['slides']))} slides")
            if b.button('↑',key='up'+obj['id'],disabled=n==0):w['service'][n-1],w['service'][n]=w['service'][n],w['service'][n-1];st.rerun()
            if c.button('↓',key='down'+obj['id'],disabled=n==len(w['service'])-1):w['service'][n+1],w['service'][n]=w['service'][n],w['service'][n+1];st.rerun()
            if d.button('✕',key='remove'+obj['id']):w['service'].pop(n);st.rerun()
            with st.expander('Preview / duplicate'):
                preview(obj)
                if st.button('Duplicate item',key='dupe'+obj['id']):w['service'].insert(n+1,duplicate(obj));st.rerun()

    with st.expander('Save / open this preparation'):
        saved_service_controls(w,st.session_state.bibles,'planner')

elif page=='Media & presentations':
    st.write('Build your image and video libraries. Use media as service items or as backgrounds behind lyrics.')
    st.caption('Upload Images.zip or Videos.zip into the matching library below. Keep your ZIPs for a full media backup; saved services include the media selected for that service.')
    images_tab,videos_tab,presentations_tab=st.tabs(['Image library','Video library','Presentations'])
    with images_tab:media_library('image')
    with videos_tab:media_library('video')
    with presentations_tab:
        st.write('Import PDF slides or the editable text from a PowerPoint presentation.')
        st.caption('For faithful PowerPoint layout, export to PDF first. PPTX import here reads text only; it does not reproduce graphics, animations or slide design.')
        upload=st.file_uploader('Choose media or presentation',type=['pdf','pptx'])
        if upload:
            st.caption(f'{upload.name} · {upload.size/1024/1024:.1f} MB')
            title=st.text_input('Media title',value=Path(upload.name).stem)
            if st.button('Import into library',type='primary'):
                try:
                    if upload.size>20*1024*1024:raise ValueError('Each upload is limited to 20 MB in this release.')
                    suffix=Path(upload.name).suffix.lower();raw=upload.getvalue()
                    if suffix=='.pdf':
                        import fitz
                        with fitz.open(stream=raw,filetype='pdf') as pdf:
                            if len(pdf)>80:raise ValueError('Limit each PDF to 80 pages.')
                            obj=item(title,'presentation','temporary')
                            obj['slides']=[{'label':f'Slide {i+1}','text':'','src':media_data(p.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).tobytes('png'),'image/png')} for i,p in enumerate(pdf)]
                    elif suffix=='.pptx':
                        from pptx import Presentation
                        prs=Presentation(io.BytesIO(raw));obj=item(title,'presentation','temporary')
                        obj['slides']=[{'label':f'Slide {i+1}','text':'\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)} for i,slide in enumerate(prs.slides)]
                        obj['slides']=[s for s in obj['slides'] if s['text'].strip()]
                        if not obj['slides']:raise ValueError('No text found. Export the slides to PDF and import that instead.')
                    else:
                        mime={'.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg','.webp':'image/webp','.mp4':'video/mp4','.webm':'video/webm'}[suffix]
                        obj=item(title,'video' if suffix in ('.mp4','.webm') else 'image',src=media_data(raw,mime))
                    candidate=copy.deepcopy(w);candidate['library'].append(obj)
                    if len(backup_bytes(candidate))>MAX_BACKUP:raise ValueError('Workspace would exceed 45 MB. Use smaller media or split services.')
                    w['library'].append(obj);notify('Imported. Find the item in Songs & library and add it to the service.')
                except Exception as exc:st.error(f'Could not import: {exc}')

elif page=='Presentation desk':
    st.subheader('Live presentation desk')
    background_controls('live')
    st.write('Select an item in the order of service below. The left screen is your preview; press **Go live** to send it to the congregation screen on the right.')
    st.caption('Click Open projector in the live desk to open a separate window. Move it to your extended second display, then click Full screen. Keep this desk open while presenting.')
    if w['service']:
        current_signature=json.dumps({'service':w['service'],'theme':w['theme'],'title':w['service_title'],'bibles':[b.get('id',b['name']) for b in st.session_state.bibles]},sort_keys=True)
        if st.button('Load / refresh live desk',type='primary') or 'live_desk_html' not in st.session_state or st.session_state.get('live_desk_version')!=6:
            st.session_state.live_desk_html=export_html(w,st.session_state.bibles,output_session=st.session_state.presenter_session).decode('utf-8')
            st.session_state.live_desk_signature=current_signature
            st.session_state.live_desk_version=6
        if st.session_state.get('live_desk_signature')!=current_signature:
            st.info('Your service or Bible versions have changed. Click Load / refresh live desk when ready to use the updated service.')
        components.html(st.session_state.live_desk_html,height=1050,scrolling=True)
        with st.expander('Save / open this prepared presentation'):
            saved_service_controls(w,st.session_state.bibles,'live')
        st.subheader('Offline presentation')
        st.download_button('Download offline presentation',data=export_html(w, st.session_state.bibles),file_name='TCC_Worship_Live.html',mime='text/html')
        st.caption('For an offline service, download and open this file in Chrome or Edge. It contains the same preview, live output and projector controls.')
    else:
        st.info('Your service is empty. Add songs, scripture or media in Service planner first.')
    with st.expander('Optional FreeShow local controls — experimental'):
        st.write('Requires this Streamlit app and FreeShow to run on the same computer. Enable the REST API in FreeShow Connections. These controls have not been tested against a running FreeShow installation.')
        if secret('TCC_LOCAL_MODE')!='1':st.info('Local controls are disabled. Start locally with TCC_LOCAL_MODE=1 to enable them.')
        else:
            port=st.number_input('FreeShow REST port',min_value=1024,max_value=65535,value=5506)
            auth=st.text_input('FreeShow auth key',type='password')
            action=st.selectbox('Command',['get_shows','next_slide','previous_slide','clear_all','restore_output'])
            if st.button('Send command to local FreeShow'):
                import urllib.parse,urllib.request
                query=urllib.parse.urlencode({'action':action,'data':'{}','auth':auth})
                try:
                    with urllib.request.urlopen(f'http://127.0.0.1:{int(port)}?{query}',timeout=4) as r:
                        response=r.read(100000).decode('utf-8',errors='replace')
                    st.success('FreeShow accepted the HTTP request.');st.code(response or '(empty response)')
                except Exception:st.error('Could not complete the FreeShow command. Check the local app, API port and auth key.')

elif page=='Saved services':
    st.subheader('Prepared presentations and song lists')
    st.write('Name your preparation and click Save. Choose a saved preparation and click Open to continue with its song list and presentation setup.')
    saved_service_controls(w,st.session_state.bibles,'library')

elif page=='Settings & backups':
    st.subheader('Presentation theme')
    background_controls('settings')
    fonts=list(FONT_FAMILIES)
    w['theme']['font_family']=st.selectbox('Font type',fonts,index=fonts.index(w['theme'].get('font_family','Arial')))
    a,b,c=st.columns(3)
    w['theme']['background']=a.color_picker('Background',w['theme']['background'])
    w['theme']['accent']=b.color_picker('Accent',w['theme']['accent'])
    w['theme']['font_size']=c.slider('Font size',24,90,w['theme']['font_size'])
    st.caption('Applies to lyrics, scripture and text slides. Long text shrinks to fit. In Presentation desk, use the font controls for immediate projector updates; save the prepared presentation there to keep live changes.')
    preview(item('Font preview','announcement','Let us worship together.'),height=240,show_background=True)
    st.subheader('Save and restore')
    st.write('Download the workspace JSON and keep your media ZIP files with it. After a server restart, restore the JSON and re-upload the ZIPs to reconnect the full collection. Saved services and offline presentations embed their selected media.')
    st.download_button('Download workspace backup',backup_bytes(w),file_name='TCC_Worship_Backup.json',mime='application/json')
    restored=st.file_uploader('Restore TCC workspace JSON',type=['json'],key='restore')
    confirm=st.checkbox('Replace the current workspace with this backup')
    if st.button('Restore backup',disabled=not(restored and confirm)):
        try:
            replacement=load_backup(restored.getvalue())
            st.session_state.workspace=replacement
            notify('Workspace restored.')
        except ValueError as exc:st.error(str(exc))
    st.subheader('Additional Bible versions')
    st.info('Open Bible desk → Browse free Bible versions to add translations by language. Download Bible JSON to retain added versions across sessions.')
    st.caption('WEB is included offline. Import additional versions that you are authorised to use. Bible imports and verse favourites are session-only; save your Bible source files separately.')
    st.code('{"name":"Version name","credit":"Required attribution","books":{"John":{"3":{"16":"Verse text"}}}}',language='json')
    bf=st.file_uploader('Import authorised Bible JSON',type=['json'],key='bible_import')
    if st.button('Add Bible version',disabled=bf is None):
        try:
            if bf.size>20*1024*1024:raise ValueError('Bible import exceeds 20 MB.')
            bible=validate_bible(json.loads(bf.getvalue()));st.session_state.bibles.append(bible);notify('Bible version added.')
        except (ValueError,TypeError) as exc:st.error(str(exc))
    st.subheader('Access and deployment')
    st.write('For a shared deployment, set APP_PASSWORD in Streamlit Secrets. Keep passwords and media backups out of a public GitHub repository. See README.md for local and cloud setup.')
    st.caption('TCC Worship v0.1 • Original implementation inspired by worship presentation workflows. No Holyrics, FreeShow, Quelea or OpenLP source code is included.')
